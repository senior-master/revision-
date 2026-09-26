"""
merge.py — Merges all project files into one final document.

USAGE:
  python3 merge.py

INSTRUCTIONS:
  1. Place this script in the same folder as all your chapter files.
  2. Make sure all files are named exactly as listed in FILES below.
  3. Run: python3 merge.py
  4. Output will be saved as: final_project.docx

NOTE: Each file starts on a new page in the merged document.
"""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy
import os

# ── FILE ORDER — edit only this list ──────────────────────────────────────────
FILES = [
    "preliminary_pages.docx",
    "chapter1.docx",
    "chapter2.docx",
    "chapter3.docx",
    "chapter4.docx",
    "chapter5.docx",
    "references_and_questionnaire.docx",
]
# ──────────────────────────────────────────────────────────────────────────────

OUTPUT = "final_project.docx"


def add_page_break(doc):
    para = doc.add_paragraph()
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def merge_documents(files, output):
    if not files:
        print("No files specified.")
        return

    missing = [f for f in files if not os.path.exists(f)]
    if missing:
        print(f"\n❌ ERROR: These files were not found:\n")
        for f in missing:
            print(f"   - {f}")
        print("\nMake sure all files are in the same folder as merge.py\n")
        return

    print(f"\n📄 Merging {len(files)} files...\n")

    merged = Document(files[0])
    print(f"  ✔ {files[0]}")

    for filepath in files[1:]:
        add_page_break(merged)
        src = Document(filepath)
        for element in src.element.body:
            if element.tag == qn("w:sectPr"):
                continue
            merged.element.body.append(deepcopy(element))
        print(f"  ✔ {filepath}")

    merged.save(output)
    print(f"\n✅ Done! Final project saved as: {output}\n")


if __name__ == "__main__":
    merge_documents(FILES, OUTPUT)
